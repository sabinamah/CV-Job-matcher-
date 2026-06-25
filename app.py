import os, json, io, re, textwrap
import fitz
from docx import Document as DocxDocument
import streamlit as st
from google import genai
from google.genai import types
from dotenv import load_dotenv

load_dotenv()
api_key_env = os.getenv("GOOGLE_API_KEY", "")

# ── CV text extractor ──────────────────────────────────────────────────────────
def extract_cv_text(file_bytes: bytes, filename: str) -> str:
    ext = filename.rsplit(".", 1)[-1].lower()
    if ext == "pdf":
        doc = fitz.open(stream=file_bytes, filetype="pdf")
        return "\n".join(page.get_text() for page in doc)
    elif ext in ("docx", "doc"):
        doc = DocxDocument(io.BytesIO(file_bytes))
        lines = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(c.text.strip() for c in row.cells if c.text.strip())
                if row_text:
                    lines.append(row_text)
        return "\n".join(lines)
    elif ext == "txt":
        return file_bytes.decode("utf-8", errors="ignore")
    return ""


# ── Gemini call ───────────────────────────────────────────────────────────────
def call_gemini(api_key: str, system: str, user_prompt: str) -> str:
    client = genai.Client(api_key=api_key)
    response = client.models.generate_content(
        model="gemini-2.5-flash",
        contents=[types.Part.from_text(text=user_prompt)],
        config=types.GenerateContentConfig(system_instruction=system),
    )
    return response.text.strip()


def call_gemini_json(api_key: str, system: str, user_prompt: str) -> dict | list:
    client = genai.Client(api_key=api_key)
    response = client.models.generate_content(
        model="gemini-2.5-flash",
        contents=[types.Part.from_text(text=user_prompt)],
        config=types.GenerateContentConfig(
            system_instruction=system,
            response_mime_type="application/json",
        ),
    )
    raw = response.text.strip().lstrip("```json").lstrip("```").rstrip("```").strip()
    return json.loads(raw)


# ── Analysis functions ─────────────────────────────────────────────────────────
ANALYSIS_SYSTEM = """You are a senior HR director and talent acquisition expert with 20+ years of experience.
You analyze CVs against job descriptions with the eye of a professional recruiter who makes hiring decisions in under 60 seconds.
You spot exaggerations, red flags, and genuine strengths instantly.
Always respond ONLY with valid JSON — no markdown, no extra text."""

def analyze_match(api_key: str, cv_text: str, job_text: str) -> dict:
    prompt = f"""Analyze this CV against the job description as an experienced HR director would in 60 seconds.

=== CV ===
{cv_text}

=== JOB DESCRIPTION ===
{job_text}

Return this exact JSON:
{{
  "match_score": <integer 0-100>,
  "verdict": "<one of: Strong Match | Good Match | Partial Match | Weak Match | Not Suitable>",
  "verdict_reason": "<2 sentences — the gut-feel HR verdict in plain language>",
  "strengths": ["<strength 1>", "<strength 2>", "<strength 3>"],
  "skill_gaps": ["<missing skill 1>", "<missing skill 2>", "<missing skill 3>"],
  "experience_match": "<Exceeds | Meets | Below | Far Below> requirements",
  "authenticity": {{
    "score": <integer 0-100 where 100 = very likely genuine>,
    "verdict": "<Likely Genuine | Suspicious | Likely Fabricated>",
    "flags": ["<red flag 1 if any>"],
    "positive_signals": ["<signal that adds credibility>"]
  }},
  "standout_skills": ["<skill that makes candidate stand out>"],
  "concerns": ["<concern an HR would have>"],
  "hire_recommendation": "<one of: Strongly Recommend | Recommend | Consider | Do Not Recommend>",
  "salary_fit": "<Likely Overpaid | Likely Fair | Likely Underpaid | Cannot Assess>",
  "years_experience_cv": "<number or range as string>",
  "years_experience_required": "<number or range from JD as string>",
  "education_match": "<Exceeds | Meets | Below | Not Required | Cannot Assess>",
  "summary_for_hiring_manager": "<3 sentences — what a hiring manager needs to know before the interview>"
}}"""
    return call_gemini_json(api_key, ANALYSIS_SYSTEM, prompt)


def generate_interview_questions(api_key: str, cv_text: str, job_text: str, match_data: dict) -> list[dict]:
    prompt = f"""You are an experienced interviewer. Based on this CV and job description, generate 8 targeted interview questions.
Mix behavioral, technical, and situational questions. Include questions that probe the identified skill gaps.

CV summary: {match_data.get('verdict_reason', '')}
Skill gaps: {', '.join(match_data.get('skill_gaps', []))}
Concerns: {', '.join(match_data.get('concerns', []))}

=== CV ===
{cv_text[:3000]}

=== JOB DESCRIPTION ===
{job_text[:2000]}

Return a JSON array of 8 objects:
[
  {{
    "question": "<the interview question>",
    "type": "<Behavioral | Technical | Situational | Culture Fit>",
    "why": "<why this question matters for this candidate — 1 sentence>",
    "what_good_answer_looks_like": "<1-2 sentences on what a strong answer includes>"
  }}
]"""
    result = call_gemini_json(api_key, ANALYSIS_SYSTEM, prompt)
    return result if isinstance(result, list) else result.get("questions", [])


def generate_project_suggestions(api_key: str, cv_text: str, job_text: str, skill_gaps: list[str]) -> list[dict]:
    gaps_str = ", ".join(skill_gaps) if skill_gaps else "general skill improvement"
    prompt = f"""You are a career coach and senior engineer. The candidate has these skill gaps for the job they're applying to: {gaps_str}

Suggest 5 concrete projects or learning activities the candidate can do to close these gaps and dramatically improve their chances.
Make them specific, achievable in 2-8 weeks, and directly relevant to the job.

=== JOB DESCRIPTION ===
{job_text[:1500]}

Return a JSON array of 5 objects:
[
  {{
    "title": "<project/activity title>",
    "description": "<what to build or do — 2-3 sentences>",
    "skills_covered": ["<skill 1>", "<skill 2>"],
    "time_estimate": "<e.g. 2 weeks>",
    "difficulty": "<Beginner | Intermediate | Advanced>",
    "how_to_showcase": "<how to add this to CV or portfolio — 1 sentence>",
    "resources": ["<free resource or platform name>"]
  }}
]"""
    result = call_gemini_json(api_key, ANALYSIS_SYSTEM, prompt)
    return result if isinstance(result, list) else []


COVER_LETTER_SYSTEM = """You are an expert career coach who writes compelling, humanized cover letters.
Your cover letters:
- Sound like a real, passionate human wrote them — NOT a robot or template
- Open with a hook that shows genuine interest, not "I am applying for..."
- Reference specific details from the job description to show research
- Tell a brief story that connects the candidate's experience to the role
- Are confident without being arrogant
- Show personality and cultural fit
- End with a clear, natural call to action
- Are 3-4 paragraphs, ~300-380 words
- Use natural language variations, occasional conversational phrasing
- Never use clichés like "team player", "go-getter", "passionate about", "leverage", "synergy"
- Feel warm, genuine, and memorable"""


def generate_cv_improvements(api_key: str, cv_text: str, job_text: str, match_data: dict) -> dict:
    skill_gaps = match_data.get("skill_gaps", [])
    concerns   = match_data.get("concerns", [])
    prompt = f"""You are a professional CV coach and career expert. Review this CV against the job description and provide specific, actionable improvement advice.

Identified skill gaps: {', '.join(skill_gaps)}
HR concerns: {', '.join(concerns)}

=== CV ===
{cv_text[:4000]}

=== JOB DESCRIPTION ===
{job_text[:2000]}

Return this exact JSON:
{{
  "overall_cv_score": <integer 0-100 rating the CV quality independent of job match>,
  "quick_wins": [
    {{"action": "<specific thing to change or add>", "why": "<why this helps — 1 sentence>", "example": "<before/after or example wording>"}}
  ],
  "structure_issues": ["<structural problem with the CV>"],
  "wording_improvements": [
    {{"original_type": "<vague phrase type e.g. 'Responsible for...' or 'Helped with...'>", "better": "<stronger alternative e.g. 'Led...', 'Delivered...'>", "why": "<why stronger>"}}
  ],
  "missing_sections": ["<section that should be added e.g. 'Quantified achievements', 'GitHub links', 'Languages'>"],
  "company_specific_tips": [
    {{"tip": "<specific thing to mention or tailor for THIS company/role>", "where_to_add": "<which CV section to add it in>"}}
  ],
  "keywords_to_add": ["<keyword from job description missing from CV>"],
  "red_flags_to_fix": ["<something on the CV that HR will flag — fix it>"],
  "length_feedback": "<Too long | Too short | About right — and why>",
  "strongest_section": "<which section of the CV is best right now>",
  "weakest_section": "<which section needs the most work>"
}}"""
    return call_gemini_json(api_key, ANALYSIS_SYSTEM, prompt)

def generate_cover_letter(api_key: str, cv_text: str, job_text: str, candidate_name: str,
                          company_name: str, tone: str, extra_notes: str) -> str:
    prompt = f"""Write a humanized cover letter for this candidate applying to this job.

Candidate name: {candidate_name or "the candidate"}
Company name: {company_name or "the company"}
Desired tone: {tone}
Extra instructions from candidate: {extra_notes or "None"}

=== CV / BACKGROUND ===
{cv_text[:3000]}

=== JOB DESCRIPTION ===
{job_text[:2000]}

Write the full cover letter. Do NOT include placeholders like [Your Name] or [Date].
Start directly with the opening paragraph. End after the sign-off.
Sign off naturally with: {candidate_name or "the candidate"}"""
    return call_gemini(api_key, COVER_LETTER_SYSTEM, prompt)


# ── Page setup ─────────────────────────────────────────────────────────────────
st.set_page_config(
    page_title="HireIQ — CV & Job Match AI",
    page_icon="🎯",
    layout="wide",
    initial_sidebar_state="collapsed"
)

st.markdown("""
<style>
@import url('https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700;800;900&display=swap');
*, html, body { font-family: 'Inter', sans-serif !important; }
.main { background: #f8f9ff !important; }
.block-container { padding: 2rem 2.5rem !important; max-width: 100% !important; }
#MainMenu, footer, header { visibility: hidden; }

.topbar {
  background: linear-gradient(135deg, #1e1b4b 0%, #3730a3 50%, #4f46e5 100%);
  padding: 28px 40px; display: flex; align-items: center; gap: 20px;
  border-radius: 20px; margin-bottom: 28px;
  box-shadow: 0 8px 32px rgba(79,70,229,0.3);
}
.topbar-title { color: white; font-size: 1.9rem; font-weight: 900; margin: 0; letter-spacing: -0.02em; }
.topbar-sub   { color: rgba(255,255,255,0.72); font-size: 0.88rem; margin: 4px 0 0; }
.pill {
  display: inline-block; background: rgba(255,255,255,0.15);
  border: 1px solid rgba(255,255,255,0.25); color: white;
  padding: 4px 14px; border-radius: 999px; font-size: 0.76rem;
  font-weight: 600; margin-top: 8px;
}

.panel {
  background: white; border-radius: 18px; border: 1px solid #e8eaf6;
  box-shadow: 0 2px 16px rgba(0,0,0,0.05); padding: 28px; margin-bottom: 20px;
}
.panel-title {
  font-size: 0.7rem; font-weight: 800; letter-spacing: 0.12em;
  text-transform: uppercase; color: #6366f1; margin-bottom: 14px;
}

/* Score ring */
.score-ring-wrap { text-align: center; padding: 10px 0 18px; }
.score-number { font-size: 4.5rem; font-weight: 900; line-height: 1; }
.score-label  { font-size: 0.85rem; color: #6b7280; margin-top: 4px; font-weight: 500; }

/* Verdict badge */
.verdict-badge {
  display: inline-flex; align-items: center; gap: 8px;
  padding: 10px 22px; border-radius: 999px;
  font-weight: 800; font-size: 1.05rem; margin: 10px 0 16px;
}

/* Tag pills */
.tag { display: inline-block; padding: 4px 12px; border-radius: 999px;
       font-size: 0.78rem; font-weight: 600; margin: 3px 3px 3px 0; }
.tag-green  { background: #d1fae5; color: #065f46; }
.tag-red    { background: #fee2e2; color: #991b1b; }
.tag-purple { background: #ede9fe; color: #5b21b6; }
.tag-blue   { background: #dbeafe; color: #1e40af; }
.tag-yellow { background: #fef3c7; color: #92400e; }
.tag-gray   { background: #f3f4f6; color: #374151; }

/* Auth card */
.auth-card {
  border-radius: 14px; padding: 18px 22px;
  border: 2px solid;
}

/* Question card */
.q-card {
  background: #fafafe; border: 1px solid #e8eaf6; border-radius: 14px;
  padding: 18px 20px; margin-bottom: 14px;
}
.q-type { font-size: 0.7rem; font-weight: 700; letter-spacing: 0.1em;
          text-transform: uppercase; margin-bottom: 6px; }
.q-text { font-size: 0.98rem; font-weight: 700; color: #1e1b4b; margin-bottom: 8px; }
.q-why  { font-size: 0.82rem; color: #6b7280; font-style: italic; }

/* Project card */
.proj-card {
  background: white; border: 1px solid #e0e7ff; border-radius: 14px;
  padding: 20px 22px; margin-bottom: 14px;
  border-left: 4px solid #6366f1;
}

/* Cover letter output */
.cl-output {
  background: #fafafe; border: 1.5px solid #c7d2fe; border-radius: 14px;
  padding: 28px 32px; font-size: 0.95rem; line-height: 1.85;
  color: #1e293b; white-space: pre-wrap;
}

/* Buttons */
.stButton > button {
  background: linear-gradient(135deg, #4f46e5, #7c3aed) !important;
  color: white !important; border: none !important;
  border-radius: 12px !important; font-weight: 700 !important;
  font-size: 0.95rem !important; padding: 0.7rem !important;
  width: 100% !important; box-shadow: 0 4px 14px rgba(79,70,229,0.35) !important;
  letter-spacing: 0.01em !important;
}
.stButton > button:disabled { opacity: 0.35 !important; }
.stTextInput > div > div > input, .stTextArea > div > div > textarea {
  border-radius: 10px !important; border: 1.5px solid #e0e7ff !important;
  font-size: 0.9rem !important; background: #fafafe !important;
}
[data-testid="stFileUploaderDropzone"] {
  background: #fafafe !important; border: 2px dashed #a5b4fc !important;
  border-radius: 12px !important;
}
.stTabs [data-baseweb="tab-list"] { gap: 8px; }
.stTabs [data-baseweb="tab"] {
  border-radius: 10px !important; font-weight: 600 !important;
  font-size: 0.88rem !important; padding: 8px 18px !important;
}
.stSelectbox > div { border-radius: 10px !important; }
hr { border: none; border-top: 1px solid #e8eaf6; margin: 20px 0; }
</style>
""", unsafe_allow_html=True)

# ── Header ─────────────────────────────────────────────────────────────────────
st.markdown("""
<div class="topbar">
  <div style="font-size:3rem">🎯</div>
  <div>
    <p class="topbar-title">HireIQ</p>
    <p class="topbar-sub">AI-Powered CV &amp; Job Match Analyzer · Instant HR-style scoring in 60 seconds · Powered by Gemini</p>
    <span class="pill">Match Score · Authenticity Check · Interview Prep · Cover Letter</span>
  </div>
</div>
""", unsafe_allow_html=True)

# ── Sidebar: API key ───────────────────────────────────────────────────────────
with st.sidebar:
    st.markdown("### 🔑 API Key")
    api_key = st.text_input("Google API Key", type="password", value=api_key_env,
                            placeholder="AIza...", label_visibility="collapsed")
    st.caption("Get a free key at **aistudio.google.com/apikey**")
    st.markdown("---")
    st.markdown("### About")
    st.caption("HireIQ acts like a senior HR director reviewing your CV against a job in under 60 seconds. "
               "It scores the match, checks authenticity, suggests what to build, generates interview questions, "
               "and writes a humanized cover letter.")

# ── Input section ──────────────────────────────────────────────────────────────
col_left, col_right = st.columns([1, 1], gap="large")

with col_left:
    st.markdown('<div class="panel">', unsafe_allow_html=True)
    st.markdown('<div class="panel-title">📄 Your CV / Resume</div>', unsafe_allow_html=True)
    uploaded_cv = st.file_uploader("Upload CV", type=["pdf", "docx", "txt"],
                                    label_visibility="collapsed", key="cv_upload")
    if uploaded_cv:
        st.success(f"✓ {uploaded_cv.name} · {uploaded_cv.size/1024:.1f} KB")
    st.markdown("**Or paste your CV text:**")
    cv_text_input = st.text_area("CV text", height=160, placeholder="Paste your CV here if you prefer...",
                                  label_visibility="collapsed", key="cv_text_input")
    st.markdown('</div>', unsafe_allow_html=True)

with col_right:
    st.markdown('<div class="panel">', unsafe_allow_html=True)
    st.markdown('<div class="panel-title">💼 Job Description</div>', unsafe_allow_html=True)
    job_text_input = st.text_area("Job description", height=200,
                                   placeholder="Paste the full job description here — the more detail, the better the analysis...",
                                   label_visibility="collapsed", key="job_text_input")
    st.markdown('</div>', unsafe_allow_html=True)

# ── Cover letter extra inputs (always visible) ─────────────────────────────────
with st.expander("✉️ Cover Letter Options (fill before analyzing)", expanded=False):
    cl_col1, cl_col2 = st.columns(2)
    with cl_col1:
        candidate_name = st.text_input("Your full name", placeholder="e.g. Sarah Johnson")
        company_name   = st.text_input("Company name",   placeholder="e.g. Google")
    with cl_col2:
        tone = st.selectbox("Cover letter tone", [
            "Professional & Confident",
            "Warm & Enthusiastic",
            "Creative & Bold",
            "Formal & Structured",
        ])
    extra_notes = st.text_area("Extra instructions (optional)",
                                placeholder="e.g. Mention my passion for sustainability. Don't mention my career gap. Emphasize leadership.",
                                height=80)

st.markdown("<br>", unsafe_allow_html=True)

# ── Analyze button ─────────────────────────────────────────────────────────────
has_cv  = uploaded_cv is not None or bool(cv_text_input.strip())
has_job = bool(job_text_input.strip())

btn_col, _ = st.columns([1, 2])
with btn_col:
    go = st.button("🎯  Analyze Match", disabled=not (has_cv and has_job and api_key))

if not api_key:
    st.warning("Enter your Google API key in the sidebar to get started.")
elif not has_cv:
    st.info("Upload your CV or paste your CV text above.")
elif not has_job:
    st.info("Paste the job description above.")

# ── Run analysis ───────────────────────────────────────────────────────────────
if go:
    # Extract CV text
    if uploaded_cv:
        cv_text = extract_cv_text(uploaded_cv.read(), uploaded_cv.name)
        if not cv_text.strip() and cv_text_input.strip():
            cv_text = cv_text_input.strip()
    else:
        cv_text = cv_text_input.strip()

    job_text = job_text_input.strip()

    if not cv_text:
        st.error("Could not extract text from your CV. Please paste it as text.")
        st.stop()

    with st.spinner("Analyzing your CV against the job — thinking like an HR director..."):
        try:
            match_data = analyze_match(api_key, cv_text, job_text)
            st.session_state["match"]    = match_data
            st.session_state["cv_text_parsed"]  = cv_text
            st.session_state["job_text"] = job_text
            st.session_state["cl_opts"]  = {
                "candidate_name": candidate_name,
                "company_name":   company_name,
                "tone":           tone,
                "extra_notes":    extra_notes,
            }
            # Clear previous lazy-loaded data
            for k in ["questions", "projects", "cover_letter"]:
                st.session_state.pop(k, None)
        except Exception as e:
            st.error(f"Analysis failed: {e}")
            st.stop()

# ── Results ────────────────────────────────────────────────────────────────────
if "match" in st.session_state:
    r   = st.session_state["match"]
    cv  = st.session_state["cv_text_parsed"]
    jd  = st.session_state["job_text"]
    cl_opts = st.session_state.get("cl_opts", {})

    score   = r.get("match_score", 0)
    verdict = r.get("verdict", "")

    VERDICT_STYLE = {
        "Strong Match":    ("#d1fae5", "#065f46", "✅"),
        "Good Match":      ("#dbeafe", "#1e40af", "👍"),
        "Partial Match":   ("#fef3c7", "#92400e", "⚠️"),
        "Weak Match":      ("#fee2e2", "#991b1b", "❌"),
        "Not Suitable":    ("#fce7f3", "#9d174d", "🚫"),
    }
    vs = VERDICT_STYLE.get(verdict, ("#f3f4f6", "#374151", "❓"))
    score_color = "#16a34a" if score >= 70 else "#d97706" if score >= 50 else "#dc2626"

    st.markdown("---")

    # ── Top summary row ────────────────────────────────────────────────────────
    top1, top2, top3, top4 = st.columns([1, 1.5, 1, 1])

    with top1:
        st.markdown('<div class="panel score-ring-wrap">', unsafe_allow_html=True)
        st.markdown(f'<div class="score-number" style="color:{score_color}">{score}</div>', unsafe_allow_html=True)
        st.markdown('<div class="score-label">Match Score / 100</div>', unsafe_allow_html=True)
        st.markdown(f"""
        <div style="margin-top:10px">
          <div style="background:#e5e7eb;border-radius:999px;height:8px;">
            <div style="background:{score_color};width:{score}%;height:8px;border-radius:999px"></div>
          </div>
        </div>""", unsafe_allow_html=True)
        st.markdown('</div>', unsafe_allow_html=True)

    with top2:
        st.markdown('<div class="panel" style="height:100%">', unsafe_allow_html=True)
        st.markdown(f'<span class="verdict-badge" style="background:{vs[0]};color:{vs[1]}">{vs[2]} {verdict}</span>', unsafe_allow_html=True)
        st.markdown(f'<p style="font-size:0.88rem;color:#374151;line-height:1.6;margin:0">{r.get("verdict_reason","")}</p>', unsafe_allow_html=True)
        st.markdown('</div>', unsafe_allow_html=True)

    with top3:
        st.markdown('<div class="panel" style="text-align:center;height:100%">', unsafe_allow_html=True)
        hire_colors = {
            "Strongly Recommend": "#16a34a",
            "Recommend":          "#2563eb",
            "Consider":           "#d97706",
            "Do Not Recommend":   "#dc2626",
        }
        hire = r.get("hire_recommendation", "")
        hcolor = hire_colors.get(hire, "#6b7280")
        st.markdown(f'<div style="font-size:0.7rem;font-weight:700;color:#9ca3af;text-transform:uppercase;letter-spacing:0.1em;margin-bottom:8px">HR Recommendation</div>', unsafe_allow_html=True)
        st.markdown(f'<div style="font-size:1rem;font-weight:800;color:{hcolor}">{hire}</div>', unsafe_allow_html=True)
        st.markdown(f'<div style="font-size:0.78rem;color:#6b7280;margin-top:10px">Experience: <strong>{r.get("experience_match","")}</strong></div>', unsafe_allow_html=True)
        st.markdown(f'<div style="font-size:0.78rem;color:#6b7280">Education: <strong>{r.get("education_match","")}</strong></div>', unsafe_allow_html=True)
        st.markdown(f'<div style="font-size:0.78rem;color:#6b7280">Salary fit: <strong>{r.get("salary_fit","")}</strong></div>', unsafe_allow_html=True)
        st.markdown('</div>', unsafe_allow_html=True)

    with top4:
        auth = r.get("authenticity", {})
        auth_score = auth.get("score", 0)
        auth_verdict = auth.get("verdict", "")
        auth_color_map = {
            "Likely Genuine":    ("#d1fae5", "#065f46"),
            "Suspicious":        ("#fef3c7", "#92400e"),
            "Likely Fabricated": ("#fee2e2", "#991b1b"),
        }
        ac = auth_color_map.get(auth_verdict, ("#f3f4f6", "#374151"))
        st.markdown(f'<div class="auth-card" style="background:{ac[0]};border-color:{ac[1]}50">', unsafe_allow_html=True)
        st.markdown(f'<div style="font-size:0.7rem;font-weight:700;color:{ac[1]};text-transform:uppercase;letter-spacing:0.1em;margin-bottom:8px">CV Authenticity</div>', unsafe_allow_html=True)
        st.markdown(f'<div style="font-size:1.15rem;font-weight:800;color:{ac[1]}">{auth_verdict}</div>', unsafe_allow_html=True)
        st.markdown(f'<div style="font-size:0.75rem;color:{ac[1]};margin-top:4px">Confidence: {auth_score}%</div>', unsafe_allow_html=True)
        if auth.get("flags"):
            for f in auth["flags"][:2]:
                st.markdown(f'<div style="font-size:0.77rem;color:#991b1b;margin-top:6px">🚩 {f}</div>', unsafe_allow_html=True)
        if auth.get("positive_signals"):
            for s in auth["positive_signals"][:1]:
                st.markdown(f'<div style="font-size:0.77rem;color:#065f46;margin-top:4px">✓ {s}</div>', unsafe_allow_html=True)
        st.markdown('</div>', unsafe_allow_html=True)

    st.markdown("<br>", unsafe_allow_html=True)

    # ── Tabs ───────────────────────────────────────────────────────────────────
    tab1, tab2, tab3, tab4, tab5 = st.tabs([
        "📊 Full Analysis",
        "❓ Interview Questions",
        "🚀 Skill-Up Projects",
        "📝 Improve Your CV",
        "✉️ Cover Letter",
    ])

    # ── Tab 1: Full Analysis ───────────────────────────────────────────────────
    with tab1:
        a1, a2 = st.columns(2, gap="large")

        with a1:
            st.markdown('<div class="panel">', unsafe_allow_html=True)
            st.markdown('<div class="panel-title">✅ Strengths</div>', unsafe_allow_html=True)
            for s in r.get("strengths", []):
                st.markdown(f'<span class="tag tag-green">✓ {s}</span>', unsafe_allow_html=True)
            st.markdown("</div>", unsafe_allow_html=True)

            st.markdown('<div class="panel">', unsafe_allow_html=True)
            st.markdown('<div class="panel-title">⭐ Standout Skills</div>', unsafe_allow_html=True)
            for s in r.get("standout_skills", []):
                st.markdown(f'<span class="tag tag-purple">⭐ {s}</span>', unsafe_allow_html=True)
            st.markdown("</div>", unsafe_allow_html=True)

            st.markdown('<div class="panel">', unsafe_allow_html=True)
            st.markdown('<div class="panel-title">📊 Experience & Education</div>', unsafe_allow_html=True)
            yrs_cv  = r.get("years_experience_cv", "N/A")
            yrs_req = r.get("years_experience_required", "N/A")
            st.markdown(f"""
            <div style="display:grid;grid-template-columns:1fr 1fr;gap:12px;margin-top:4px">
              <div style="background:#f8f9ff;border-radius:10px;padding:14px;text-align:center">
                <div style="font-size:1.5rem;font-weight:800;color:#4f46e5">{yrs_cv}</div>
                <div style="font-size:0.72rem;color:#6b7280;font-weight:600;margin-top:2px">YEARS IN CV</div>
              </div>
              <div style="background:#f8f9ff;border-radius:10px;padding:14px;text-align:center">
                <div style="font-size:1.5rem;font-weight:800;color:#6366f1">{yrs_req}</div>
                <div style="font-size:0.72rem;color:#6b7280;font-weight:600;margin-top:2px">YEARS REQUIRED</div>
              </div>
            </div>
            """, unsafe_allow_html=True)
            st.markdown("</div>", unsafe_allow_html=True)

        with a2:
            st.markdown('<div class="panel">', unsafe_allow_html=True)
            st.markdown('<div class="panel-title">❌ Skill Gaps</div>', unsafe_allow_html=True)
            for s in r.get("skill_gaps", []):
                st.markdown(f'<span class="tag tag-red">✗ {s}</span>', unsafe_allow_html=True)
            st.markdown("</div>", unsafe_allow_html=True)

            st.markdown('<div class="panel">', unsafe_allow_html=True)
            st.markdown('<div class="panel-title">⚠️ HR Concerns</div>', unsafe_allow_html=True)
            for c in r.get("concerns", []):
                st.markdown(f'<span class="tag tag-yellow">⚠ {c}</span>', unsafe_allow_html=True)
            if not r.get("concerns"):
                st.markdown('<span class="tag tag-green">No major concerns</span>', unsafe_allow_html=True)
            st.markdown("</div>", unsafe_allow_html=True)

            st.markdown('<div class="panel">', unsafe_allow_html=True)
            st.markdown('<div class="panel-title">📝 Summary for Hiring Manager</div>', unsafe_allow_html=True)
            st.markdown(f'<p style="font-size:0.9rem;color:#374151;line-height:1.7;margin:0">{r.get("summary_for_hiring_manager","")}</p>', unsafe_allow_html=True)
            st.markdown("</div>", unsafe_allow_html=True)

    # ── Tab 2: Interview Questions ─────────────────────────────────────────────
    with tab2:
        if "questions" not in st.session_state:
            if st.button("Generate Interview Questions", key="gen_q"):
                with st.spinner("Generating targeted interview questions..."):
                    try:
                        qs = generate_interview_questions(api_key, cv, jd, r)
                        st.session_state["questions"] = qs
                        st.rerun()
                    except Exception as e:
                        st.error(f"Failed: {e}")
        else:
            qs = st.session_state["questions"]
            type_colors = {
                "Behavioral": "#dbeafe",
                "Technical":  "#ede9fe",
                "Situational": "#fef3c7",
                "Culture Fit": "#d1fae5",
            }
            for i, q in enumerate(qs, 1):
                qtype = q.get("type", "General")
                tc = type_colors.get(qtype, "#f3f4f6")
                st.markdown(f"""
                <div class="q-card">
                  <div class="q-type" style="color:#4f46e5">
                    <span style="background:{tc};padding:2px 10px;border-radius:999px">{qtype}</span>
                    &nbsp; Q{i}
                  </div>
                  <div class="q-text">{q.get('question','')}</div>
                  <div class="q-why">💡 Why: {q.get('why','')}</div>
                  <div style="font-size:0.8rem;color:#374151;margin-top:8px;background:#f0f4ff;padding:8px 12px;border-radius:8px">
                    ✅ Strong answer includes: {q.get('what_good_answer_looks_like','')}
                  </div>
                </div>
                """, unsafe_allow_html=True)

            if st.button("↻ Regenerate Questions", key="regen_q"):
                st.session_state.pop("questions", None)
                st.rerun()

    # ── Tab 3: Skill-Up Projects ───────────────────────────────────────────────
    with tab3:
        gaps = r.get("skill_gaps", [])
        if gaps:
            st.markdown(f'<p style="color:#6b7280;font-size:0.88rem">Based on your skill gaps: {", ".join(f"<strong>{g}</strong>" for g in gaps)} — here are concrete projects to close them.</p>', unsafe_allow_html=True)

        if "projects" not in st.session_state:
            if st.button("Generate Project Suggestions", key="gen_proj"):
                with st.spinner("Thinking of projects to make you hireable..."):
                    try:
                        projs = generate_project_suggestions(api_key, cv, jd, gaps)
                        st.session_state["projects"] = projs
                        st.rerun()
                    except Exception as e:
                        st.error(f"Failed: {e}")
        else:
            projs = st.session_state["projects"]
            diff_colors = {"Beginner": "#d1fae5", "Intermediate": "#fef3c7", "Advanced": "#fee2e2"}
            for i, p in enumerate(projs, 1):
                dc = diff_colors.get(p.get("difficulty", ""), "#f3f4f6")
                st.markdown(f"""
                <div class="proj-card">
                  <div style="display:flex;justify-content:space-between;align-items:flex-start;gap:10px">
                    <div style="font-size:1rem;font-weight:800;color:#1e1b4b">#{i} {p.get('title','')}</div>
                    <div style="display:flex;gap:6px;flex-shrink:0">
                      <span class="tag" style="background:{dc};color:#374151">{p.get('difficulty','')}</span>
                      <span class="tag tag-blue">⏱ {p.get('time_estimate','')}</span>
                    </div>
                  </div>
                  <p style="font-size:0.88rem;color:#374151;margin:10px 0 10px;line-height:1.6">{p.get('description','')}</p>
                  <div style="margin-bottom:8px">
                    {"".join(f'<span class="tag tag-purple">{s}</span>' for s in p.get('skills_covered', []))}
                  </div>
                  <div style="font-size:0.8rem;color:#16a34a;background:#f0fdf4;padding:8px 12px;border-radius:8px;margin-bottom:8px">
                    📌 Portfolio: {p.get('how_to_showcase','')}
                  </div>
                  <div style="font-size:0.78rem;color:#6b7280">
                    Resources: {" · ".join(p.get('resources', []))}
                  </div>
                </div>
                """, unsafe_allow_html=True)

            if st.button("↻ Regenerate Suggestions", key="regen_proj"):
                st.session_state.pop("projects", None)
                st.rerun()

    # ── Tab 4: Improve Your CV ────────────────────────────────────────────────
    with tab4:
        st.markdown('<p style="color:#6b7280;font-size:0.88rem">Get specific advice to improve your CV for this role — wording, structure, keywords, and company-specific tailoring.</p>', unsafe_allow_html=True)
        if "cv_improvements" not in st.session_state:
            if st.button("📝 Analyze & Improve My CV", key="gen_imp"):
                with st.spinner("Reviewing your CV in detail..."):
                    try:
                        imp = generate_cv_improvements(api_key, cv, jd, r)
                        st.session_state["cv_improvements"] = imp
                        st.rerun()
                    except Exception as e:
                        st.error(f"Failed: {e}")
        else:
            imp = st.session_state["cv_improvements"]

            # Overall CV score
            cv_score = imp.get("overall_cv_score", 0)
            cv_score_color = "#16a34a" if cv_score >= 70 else "#d97706" if cv_score >= 50 else "#dc2626"
            sc1, sc2, sc3 = st.columns(3)
            with sc1:
                st.markdown(f'<div class="panel score-ring-wrap"><div class="score-number" style="color:{cv_score_color}">{cv_score}</div><div class="score-label">CV Quality Score</div><div style="margin-top:10px"><div style="background:#e5e7eb;border-radius:999px;height:8px"><div style="background:{cv_score_color};width:{cv_score}%;height:8px;border-radius:999px"></div></div></div></div>', unsafe_allow_html=True)
            with sc2:
                st.markdown(f'<div class="panel" style="text-align:center;padding:20px"><div style="font-size:0.7rem;font-weight:700;color:#9ca3af;text-transform:uppercase;letter-spacing:0.1em;margin-bottom:8px">Strongest Section</div><div style="font-size:0.95rem;font-weight:800;color:#16a34a">{imp.get("strongest_section","")}</div></div>', unsafe_allow_html=True)
            with sc3:
                st.markdown(f'<div class="panel" style="text-align:center;padding:20px"><div style="font-size:0.7rem;font-weight:700;color:#9ca3af;text-transform:uppercase;letter-spacing:0.1em;margin-bottom:8px">Weakest Section</div><div style="font-size:0.95rem;font-weight:800;color:#dc2626">{imp.get("weakest_section","")}</div></div>', unsafe_allow_html=True)

            st.markdown("<br>", unsafe_allow_html=True)
            imp1, imp2 = st.columns(2, gap="large")

            with imp1:
                # Quick wins
                if imp.get("quick_wins"):
                    st.markdown('<div class="panel"><div class="panel-title">⚡ Quick Wins</div>', unsafe_allow_html=True)
                    for qw in imp["quick_wins"]:

                        example_html = f'<div style="font-size:0.78rem;color:#6b7280;margin-top:6px;font-style:italic">e.g. {qw["example"]}</div>' if qw.get("example") else ""
                        st.markdown(f'...{example_html}...', unsafe_allow_html=True)

                # Company-specific tips
                if imp.get("company_specific_tips"):
                    st.markdown('<div class="panel"><div class="panel-title">🏢 Company-Specific Tailoring</div>', unsafe_allow_html=True)
                    for tip in imp["company_specific_tips"]:
                        st.markdown(f'<div style="background:#f0f4ff;border-left:3px solid #6366f1;border-radius:0 10px 10px 0;padding:12px 14px;margin-bottom:10px"><div style="font-size:0.88rem;font-weight:700;color:#3730a3">{tip.get("tip","")}</div><div style="font-size:0.78rem;color:#6b7280;margin-top:4px">Add to: <strong>{tip.get("where_to_add","")}</strong></div></div>', unsafe_allow_html=True)
                    st.markdown("</div>", unsafe_allow_html=True)

                # Missing keywords
                if imp.get("keywords_to_add"):
                    st.markdown('<div class="panel"><div class="panel-title">🔑 Keywords to Add</div>', unsafe_allow_html=True)
                    for kw in imp["keywords_to_add"]:
                        st.markdown(f'<span class="tag tag-blue">+ {kw}</span>', unsafe_allow_html=True)
                    st.markdown("</div>", unsafe_allow_html=True)

            with imp2:
                # Wording improvements
                if imp.get("wording_improvements"):
                    st.markdown('<div class="panel"><div class="panel-title">✍️ Stronger Wording</div>', unsafe_allow_html=True)
                    for wi in imp["wording_improvements"]:
                        st.markdown(f'<div style="margin-bottom:12px"><div style="font-size:0.78rem;color:#991b1b;text-decoration:line-through">{wi.get("original_type","")}</div><div style="font-size:0.88rem;font-weight:700;color:#065f46;margin-top:2px">→ {wi.get("better","")}</div><div style="font-size:0.75rem;color:#6b7280;margin-top:2px">{wi.get("why","")}</div></div>', unsafe_allow_html=True)
                    st.markdown("</div>", unsafe_allow_html=True)

                # Missing sections
                if imp.get("missing_sections"):
                    st.markdown('<div class="panel"><div class="panel-title">➕ Missing Sections</div>', unsafe_allow_html=True)
                    for ms in imp["missing_sections"]:
                        st.markdown(f'<span class="tag tag-yellow">+ {ms}</span>', unsafe_allow_html=True)
                    st.markdown("</div>", unsafe_allow_html=True)

                # Red flags to fix
                if imp.get("red_flags_to_fix"):
                    st.markdown('<div class="panel"><div class="panel-title">🚩 Red Flags to Fix</div>', unsafe_allow_html=True)
                    for rf in imp["red_flags_to_fix"]:
                        st.markdown(f'<div style="background:#fef2f2;border-left:3px solid #dc2626;border-radius:0 10px 10px 0;padding:10px 14px;margin-bottom:8px;font-size:0.85rem;color:#991b1b">{rf}</div>', unsafe_allow_html=True)
                    st.markdown("</div>", unsafe_allow_html=True)

                # Length feedback
                if imp.get("length_feedback"):
                    st.markdown(f'<div class="panel"><div class="panel-title">📏 Length & Format</div><p style="font-size:0.88rem;color:#374151;margin:0">{imp["length_feedback"]}</p></div>', unsafe_allow_html=True)

            if st.button("↻ Re-analyze", key="regen_imp"):
                st.session_state.pop("cv_improvements", None)
                st.rerun()

    # ── Tab 5: Cover Letter ────────────────────────────────────────────────────
    with tab5:
        st.markdown('<p style="color:#6b7280;font-size:0.88rem">Generate a humanized, non-robotic cover letter tailored to this exact job.</p>', unsafe_allow_html=True)

        # Allow updating options inline
        tc1, tc2 = st.columns(2)
        with tc1:
            cl_name    = st.text_input("Your name",    value=cl_opts.get("candidate_name",""), key="cl_name2")
            cl_company = st.text_input("Company name", value=cl_opts.get("company_name",""),   key="cl_company2")
        with tc2:
            cl_tone = st.selectbox("Tone", [
                "Professional & Confident",
                "Warm & Enthusiastic",
                "Creative & Bold",
                "Formal & Structured",
            ], index=["Professional & Confident","Warm & Enthusiastic","Creative & Bold","Formal & Structured"].index(
                cl_opts.get("tone","Professional & Confident")), key="cl_tone2")
        cl_notes = st.text_area("Extra instructions", value=cl_opts.get("extra_notes",""),
                                 height=70, key="cl_notes2",
                                 placeholder="e.g. Emphasize my leadership. Mention my relocation to London.")

        gen_cl = st.button("✉️ Generate Cover Letter", key="gen_cl")

        if gen_cl:
            with st.spinner("Writing your cover letter — making it sound like a real human..."):
                try:
                    cl_text = generate_cover_letter(
                        api_key, cv, jd,
                        cl_name, cl_company, cl_tone, cl_notes
                    )
                    st.session_state["cover_letter"] = cl_text
                    st.rerun()
                except Exception as e:
                    st.error(f"Failed: {e}")

        if "cover_letter" in st.session_state:
            cl_text = st.session_state["cover_letter"]
            st.markdown(f'<div class="cl-output">{cl_text}</div>', unsafe_allow_html=True)
            st.markdown("<br>", unsafe_allow_html=True)

            dl_col1, dl_col2 = st.columns(2)
            with dl_col1:
                st.download_button(
                    "📥 Download Cover Letter (.txt)",
                    data=cl_text,
                    file_name=f"cover_letter_{cl_name or 'candidate'}.txt",
                    mime="text/plain",
                    use_container_width=True,
                )
            with dl_col2:
                if st.button("↻ Rewrite", key="regen_cl"):
                    st.session_state.pop("cover_letter", None)
                    st.rerun()
